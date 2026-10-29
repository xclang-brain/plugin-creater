"""申し送りの3つの形の比較の模擬ユーザーが示すファイル（docx と xlsx）と、内側の EVAL の入力を作る。
使い方：<仮想環境>/bin/python make_files.py [--out 出力先のフォルダ] [--inner 内側の EVAL の入力のフォルダ]
（出力先の既定は files/ と inner/inputs/）。部品の版は ../requirements.txt。
題材は、アイデア創出スキル（make-idea）の試行で出たアイデア（AI の申し送りを入れる前に、AI なしの良い形と並べて
3つの形を比べる）。患者、病室、記録、人の役割、数はすべて架空で、実在の患者や病院のデータではない。

形（申し送りの作り方）
  (a) AI なし：事実は電子カルテから自動で入れ、見立て（病状の重さ）は人が書く
  (b) AI が全体を下書きし、人が直す
  (c) (a) の形のまま、事実の欄だけを AI が埋める（各事実に出どころの記録と時刻を付ける）
判定者には形を伏せ、記号（X・Y・Z）だけを見せる。記号は患者ごとに付け替える。
"""
import argparse
import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font

FIXED = datetime.datetime(2026, 10, 26, 9, 0, 0)  # 作るたびに中身が変わらないよう、作成日時を固定する
AUTHOR = "看護部"

FIELDS = ["病状の重さ", "患者の要約", "すること", "気をつけること", "受け手の要約"]
COUNTED = FIELDS[:4]  # 受け手の要約は、受け手役が書く欄なので漏れに数えない
FORMS = ["a", "b", "c"]
FORM_NAMES = {"a": "AI なし", "b": "AI が全体を下書き", "c": "事実の欄だけ AI"}
LABELS = ["X", "Y", "Z"]
# 正解の一覧の項目番号と欄（どの患者も同じ並び）
ITEM_FIELDS = {1: "病状の重さ", 2: "患者の要約", 3: "患者の要約", 4: "すること", 5: "すること", 6: "気をつけること", 7: "受け手の要約"}


def patient(pid, room, items, labels, omit, disputes=None, severity=None, sev_disputes=None, noticed=None, time=None, said=None):
    """患者1人分。items は項目番号 2〜7 の文と、項目番号 1 の見立て（安定・要注意…）。
    omit：形ごとの、最終の判定（合議の後）で漏れの項目番号。disputes：{(形, 項目番号): (判定者1, 判定者2, 合議, あいまいな文)}。
    severity：形ごとの見立ての判定（同意・不同意。項目 1 が漏れなら書かれていない）。sev_disputes：{形: (判定者1, 判定者2, 合議)}。
    noticed：形ごとの受け手が気づいた欠けの数。time：形ごとの (書く分, 直す分, 備考)。said：受け手の要約に書いた文。"""
    return {
        "id": pid, "room": room, "items": items, "labels": labels, "omit": omit, "disputes": disputes or {},
        "severity": severity or {}, "sev_disputes": sev_disputes or {}, "noticed": noticed or {}, "time": time, "said": said,
    }


# ---- 前回の1回（示すファイル）：架空の患者6人。結論は (a) ----------------------------------------------------------
PREVIOUS = [
    patient("P-01", "412", {
        1: "要注意（夜間に息切れが強くなる）", 2: "COPD の悪化で入院4日目", 3: "酸素 鼻カニューラ 1L/分",
        4: "22時に吸入", 5: "6時に体温と SpO2 の測定", 6: "SpO2 が 88% を下回ったら当直医に連絡", 7: "SpO2 88% で連絡することを復唱した",
    }, {"X": "c", "Y": "a", "Z": "b"},
        omit={"a": [3, 6], "b": [4], "c": [6]},
        severity={"a": "同意", "b": "同意", "c": "同意"}, noticed={"a": 1, "b": 0, "c": 0},
        time={"a": (5, 1, ""), "b": (0, 9, ""), "c": (4, 3, "")}),
    patient("P-02", "415", {
        1: "安定", 2: "脳梗塞でリハビリ中", 3: "右手に軽いまひ",
        4: "21時に血圧の測定", 5: "夜間はトイレの付き添い", 6: "ふらつきがあるので、歩くときは必ず付き添う", 7: "トイレの付き添いを復唱した",
    }, {"X": "a", "Y": "b", "Z": "c"},
        omit={"a": [5, 7], "b": [6], "c": [2, 5]},
        severity={"a": "同意", "b": "同意", "c": "同意"}, noticed={"a": 0, "b": 0, "c": 1},
        time={"a": (6, 1, ""), "b": (0, 8, ""), "c": (4, 4, "")}),
    patient("P-03", "418", {
        1: "安定", 2: "乳がんの手術の後2日目", 3: "ドレーンが1本入っている",
        4: "ドレーンの排液の量を記録", 5: "6時に痛みの点数を記録", 6: "排液が赤く増えたら当直医に連絡", 7: "排液の確認を復唱した",
    }, {"X": "b", "Y": "c", "Z": "a"},
        omit={"a": [4, 6], "b": [3], "c": [6]},
        severity={"a": "不同意", "b": "同意", "c": "不同意"}, noticed={"a": 1, "b": 1, "c": 0},
        time={"a": (5, 1, ""), "b": (0, 10, ""), "c": (3, 3, "")}),
    patient("P-04", "421", {
        1: "要注意（熱が下がりきっていない）", 2: "腎盂腎炎で抗菌薬の点滴中", 3: "37.8℃ の発熱が続いている",
        4: "0時に抗菌薬の点滴", 5: "尿の量を記録", 6: "38.5℃ を超えたら当直医に連絡", 7: "発熱時の連絡を復唱した",
    }, {"X": "a", "Y": "c", "Z": "b"},
        omit={"a": [2], "b": [7], "c": [4]},
        severity={"a": "同意", "b": "同意", "c": "同意"}, noticed={"a": 0, "b": 0, "c": 0},
        time={"a": (6, 1, ""), "b": (0, 9, ""), "c": (4, 3, "")}),
    patient("P-05", "423", {
        1: "要注意（便に血が混じる）", 2: "消化管出血の止血の後", 3: "絶食中",
        4: "6時に採血（ヘモグロビン）", 5: "便の色を観察", 6: "黒い便や血圧の低下があれば当直医に連絡", 7: "便の観察を復唱した",
    }, {"X": "b", "Y": "a", "Z": "c"},
        omit={"a": [5, 6], "b": [1], "c": [5, 6]},
        severity={"a": "同意", "b": "書かれていない", "c": "同意"}, noticed={"a": 1, "b": 0, "c": 1},
        time={"a": (5, 1, ""), "b": (0, 9, ""), "c": (4, 4, "")}),
    patient("P-06", "426", {
        1: "安定", 2: "誤嚥性肺炎で入院6日目", 3: "とろみをつけた食事",
        4: "食事のときは体を起こす", 5: "22時に口の中の清潔", 6: "むせが続いたら食事を止める", 7: "食事の姿勢を復唱した",
    }, {"X": "c", "Y": "b", "Z": "a"},
        omit={"a": [4], "b": [5], "c": []},
        disputes={("b", 5): ("漏れ", "あり", "漏れ", "口の中のケア")},
        severity={"a": "同意", "b": "不同意", "c": "同意"}, noticed={"a": 0, "b": 0, "c": 0},
        time={"a": (6, 1, ""), "b": (0, 9, ""), "c": (3, 3, "")}),
]
PREVIOUS_EXPECTED = {  # 前回の報告の数（形ごとの欄の漏れ、見立ての不同意、1件あたりの時間、受け手が気づいた欠け）
    "a": ([0, 2, 4, 3], 1, 6.5, 3),
    "b": ([1, 1, 2, 1], 1, 9.0, 1),
    "c": ([0, 1, 3, 3], 1, 7.0, 2),
}
PREVIOUS_CONCLUSION = "a"

# ---- 内側の EVAL の入力：架空の患者6人。正しい結論は (c) ------------------------------------------------------------
# (b) は漏れと時間の条件に当てはまり、漏れの合計も最も少ないが、見立ての不同意が (a) より多いので入れない。
# (c) の P-13 は備考に「中断」があり、平均から外す（外さないと時間の増え方が2分を超える）。
# 受け手の要約の漏れ（(c) の2件）は数えない（数えると (c) の漏れの減りが3件に届かない）。
# 判定者2人の判定が分かれた項目は、合議の列の判定を使う（(a) に4件）。記号は患者ごとに付け替えてある。
INNER = [
    patient("P-11", None, {
        1: "要注意（夜間に酸素の量が増えている）", 2: "肺炎で入院3日目。抗菌薬の点滴中", 3: "酸素 鼻カニューラ 2L/分、SpO2 94%",
        4: "22時に抗菌薬の点滴", 5: "6時に採血（CRP）", 6: "SpO2 が 90% を下回ったら当直医に連絡", 7: "SpO2 90% で連絡することを復唱した",
    }, {"X": "b", "Y": "a", "Z": "c"},
        omit={"a": [3], "b": [5], "c": []},
        disputes={("a", 6): ("あり", "漏れ", "あり", "SpO2 低下時は連絡")},
        severity={"a": "同意", "b": "不同意", "c": "同意"}, noticed={"a": 1, "b": 0, "c": 0},
        time={"a": (5, 1, ""), "b": (0, 7, ""), "c": (4, 3, "")}),
    patient("P-12", None, {
        1: "安定", 2: "大腿骨の骨折の手術の後2日目。痛み止めは定時で内服", 3: "尿の管は明日抜く予定",
        4: "2時間ごとに体の向きを変える", 5: "21時に痛みの点数を記録", 6: "夜間にせん妄が出やすい。ベッドの柵を確かめる", 7: "せん妄のときの対応を復唱した",
    }, {"X": "c", "Y": "b", "Z": "a"},
        omit={"a": [5, 6], "b": [], "c": [5, 6]},
        disputes={("a", 3): ("漏れ", "あり", "あり", "尿の管あり")},
        severity={"a": "同意", "b": "不同意", "c": "同意"}, noticed={"a": 1, "b": 0, "c": 1},
        time={"a": (4, 1, ""), "b": (0, 6, ""), "c": (3, 4, "")}),
    patient("P-13", None, {
        1: "要注意（夕方に低血糖があった）", 2: "糖尿病でインスリンの量を調整中", 3: "夕食は半分食べた",
        4: "21時と6時に血糖の測定", 5: "血糖が 70 未満ならブドウ糖を飲んでもらう", 6: "低血糖の症状（冷や汗、ふるえ）に気をつける", 7: "21時の血糖の測定を復唱した",
    }, {"X": "a", "Y": "c", "Z": "b"},
        omit={"a": [3, 5], "b": [1], "c": [7]},
        disputes={("a", 5): ("あり", "漏れ", "漏れ", "低いときは対応")},
        severity={"a": "同意", "b": "書かれていない", "c": "同意"}, noticed={"a": 0, "b": 0, "c": 0},
        time={"a": (6, 1, ""), "b": (0, 8, ""), "c": (10, 8, "中断（隣の部屋の急変の対応）")}),
    patient("P-14", None, {
        1: "安定", 2: "胆のう炎で絶食中。点滴で水分", 3: "腹痛は痛み止めで落ち着いている",
        4: "0時に点滴の入れ替え", 5: "明日の朝の腹部エコーまで絶食を続ける", 6: "38℃ 以上の熱が出たら当直医に連絡", 7: "熱が出たときの連絡を復唱した",
    }, {"X": "b", "Y": "c", "Z": "a"},
        omit={"a": [4, 6], "b": [3], "c": [6]},
        severity={"a": "不同意", "b": "不同意", "c": "不同意"}, noticed={"a": 1, "b": 1, "c": 0},
        time={"a": (5, 1, ""), "b": (0, 7, ""), "c": (4, 3, "")}),
    patient("P-15", None, {
        1: "要注意（体重が2日で1.5kg 増えた）", 2: "心不全の悪化で入院5日目", 3: "利尿薬を朝に内服",
        4: "6時に体重の測定", 5: "水分の量を記録（1日 1000mL まで）", 6: "横になると息苦しいと言ったら当直医に連絡", 7: "体重の測定を復唱した",
    }, {"X": "c", "Y": "a", "Z": "b"},
        omit={"a": [3], "b": [1], "c": [4, 7]},
        disputes={("a", 5): ("漏れ", "あり", "あり", "水分に気をつける")},
        severity={"a": "同意", "b": "書かれていない", "c": "同意"}, noticed={"a": 0, "b": 0, "c": 0},
        time={"a": (6, 1, ""), "b": (0, 7, ""), "c": (4, 4, "")}),
    patient("P-16", None, {
        1: "安定", 2: "虫垂炎の手術の後1日目。歩き始めた", 3: "傷の痛みは軽い",
        4: "明日の朝から流動食を始める", 5: "22時に傷の観察", 6: "吐き気が出たら食事を止めて連絡", 7: "食事の開始を復唱した",
    }, {"X": "a", "Y": "b", "Z": "c"},
        omit={"a": [5, 6], "b": [6], "c": [3, 6]},
        severity={"a": "同意", "b": "同意", "c": "同意"}, sev_disputes={"b": ("不同意", "同意", "同意")},
        noticed={"a": 1, "b": 0, "c": 2},
        time={"a": (4, 1, ""), "b": (0, 7, ""), "c": (3, 5, "")}),
]
INNER_EXPECTED = {
    "a": ([0, 3, 4, 3], 1, 6.0, 4),
    "b": ([2, 1, 1, 1], 3, 7.0, 1),
    "c": ([0, 1, 2, 3], 1, 7.4, 3),
}
INNER_CONCLUSION = "c"


# ---- 数える（作ったデータが、ねらいの数になっているかを確かめる） ----------------------------------------------------

def summarize(patients):
    out = {}
    for form in FORMS:
        fields = [0, 0, 0, 0]
        disagree = 0
        noticed = 0
        minutes = []
        for p in patients:
            for no in p["omit"][form]:
                field = ITEM_FIELDS[no]
                if field in COUNTED:
                    fields[COUNTED.index(field)] += 1
            final_sev = p["sev_disputes"].get(form, (None, None, p["severity"][form]))[2]
            disagree += final_sev == "不同意"
            noticed += p["noticed"][form]
            write, fix, note = p["time"][form]
            if "中断" not in note:
                minutes.append(write + fix)
        out[form] = (fields, disagree, round(sum(minutes) / len(minutes), 1), noticed)
    return out


def decide(summary):
    base_total = sum(summary["a"][0])
    base_time = summary["a"][2]
    passed = []
    for form in ("b", "c"):
        total = sum(summary[form][0])
        if base_total - total >= 3 and summary[form][2] - base_time <= 2 and summary[form][1] <= summary["a"][1]:
            passed.append((total, form))
    return min(passed)[1] if passed else "a"


def check_data(patients, expected, conclusion):
    for p in patients:
        assert sorted(p["labels"].values()) == FORMS, p["id"]
        for (form, no), (j1, j2, agreed, _vague) in p["disputes"].items():
            assert j1 != j2, (p["id"], form, no)
            assert (agreed == "漏れ") == (no in p["omit"][form]), (p["id"], form, no)
        for form in FORMS:
            assert (p["severity"][form] == "書かれていない") == (1 in p["omit"][form]), (p["id"], form)
    summary = summarize(patients)
    assert summary == expected, summary
    assert decide(summary) == conclusion, decide(summary)


# ---- 作る -----------------------------------------------------------------------------------------------------------

def base_doc() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    return doc


def fix_doc(doc: Document, title: str) -> None:
    props = doc.core_properties
    props.author = AUTHOR
    props.last_modified_by = AUTHOR
    props.title = title
    props.created = FIXED
    props.modified = FIXED


def fix_book(wb: Workbook) -> None:
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = FIXED
    wb.properties.modified = FIXED


def bold_header(ws, widths) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCDEFG", widths):
        ws.column_dimensions[col].width = width


def make_reference(path: Path, patients) -> None:
    """正解の一覧（先輩の看護師が作る）：患者、項目番号、欄、書かれるべきこと"""
    wb = Workbook()
    ws = wb.active
    ws.title = "正解の一覧"
    ws.append(["患者", "項目番号", "欄", "書かれるべきこと"])
    for p in patients:
        for no in range(1, 8):
            text = p["items"][no]
            if no == 1:
                text = f"病状の重さ（見立て）が書かれている：{text}"
            elif no == 7:
                text = f"受け手が復唱する：{text.removesuffix('を復唱した')}"
            ws.append([p["id"], no, ITEM_FIELDS[no], text])
    bold_header(ws, (8, 10, 14, 60))
    fix_book(wb)
    wb.save(path)


def judged(p, form, no):
    if (form, no) in p["disputes"]:
        j1, j2, agreed, _ = p["disputes"][(form, no)]
        return j1, j2, agreed
    value = "漏れ" if no in p["omit"][form] else "あり"
    return value, value, ""


def make_judgments(path: Path, patients) -> None:
    """判定シート（判定者2人が、形を伏せた記号で付ける）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "漏れ"
    ws.append(["患者", "記号", "項目番号", "判定者1", "判定者2", "合議"])
    for p in patients:
        for label in LABELS:
            form = p["labels"][label]
            for no in range(1, 8):
                ws.append([p["id"], label, no, *judged(p, form, no)])
    bold_header(ws, (8, 6, 10, 10, 10, 10))
    sev = wb.create_sheet("見立て")
    sev.append(["患者", "記号", "判定者1", "判定者2", "合議"])
    for p in patients:
        for label in LABELS:
            form = p["labels"][label]
            if form in p["sev_disputes"]:
                sev.append([p["id"], label, *p["sev_disputes"][form]])
            else:
                sev.append([p["id"], label, p["severity"][form], p["severity"][form], ""])
    bold_header(sev, (8, 6, 14, 14, 14))
    recv = wb.create_sheet("受け手")
    recv.append(["患者", "記号", "気づいた欠け（件）"])
    for p in patients:
        for label in LABELS:
            recv.append([p["id"], label, p["noticed"][p["labels"][label]]])
    bold_header(recv, (8, 6, 18))
    legend = wb.create_sheet("凡例")
    legend.append(["項目", "意味"])
    legend.append(["記号", "申し送りの形を伏せるための記号（X・Y・Z）"])
    legend.append(["あり", "正解の一覧の項目が、申し送りに書かれていた"])
    legend.append(["漏れ", "正解の一覧の項目が、申し送りに書かれていなかった"])
    legend.append(["同意・不同意", "申し送りの病状の重さ（見立て）に、判定者が同意したか"])
    legend.append(["書かれていない", "申し送りに病状の重さが書かれていなかった"])
    legend.append(["合議", "判定者2人の判定が分かれた項目は、話し合って決めた判定を合議の列に書く。合議の列に判定が書いてあるときは、合議の列の判定を使う"])
    legend.append(["気づいた欠け", "受け手役が受け手の要約を書くときに、足りないと気づいて聞き返したことの数"])
    bold_header(legend, (14, 90))
    fix_book(wb)
    wb.save(path)


def make_timelog(path: Path, patients) -> None:
    """時間の記録（申し送りを作った人が付ける）と、記号の対応"""
    wb = Workbook()
    ws = wb.active
    ws.title = "時間"
    ws.append(["患者", "形", "書く（分）", "直す（分）", "備考"])
    for p in patients:
        for form in FORMS:
            write, fix, note = p["time"][form]
            ws.append([p["id"], f"({form})", write, fix, note])
    bold_header(ws, (8, 6, 10, 10, 32))
    key = wb.create_sheet("記号の対応")
    key.append(["患者", "記号", "形"])
    for p in patients:
        for label in LABELS:
            key.append([p["id"], label, f"({p['labels'][label]})"])
    bold_header(key, (8, 6, 6))
    note = wb.create_sheet("メモ")
    note.append(["列", "意味"])
    note.append(["書く（分）", "人が申し送りを書いた時間"])
    note.append(["直す（分）", "電子カルテや AI が入れた文を読んで、確かめて直した時間"])
    note.append(["(a)", "AI なし：事実は電子カルテから自動で入れ、見立ては人が書く"])
    note.append(["(b)", "AI が全体を下書きし、人が直す"])
    note.append(["(c)", "(a) の形のまま、事実の欄だけを AI が埋める（出どころの記録と時刻つき）"])
    bold_header(note, (12, 70))
    fix_book(wb)
    wb.save(path)


SOURCE_TIMES = {2: "08:00", 3: "17:00", 4: "15:30", 5: "16:10"}


def handoff_lines(p, form, date):
    lines = []
    sev = p["items"][1]
    if 1 not in p["omit"][form]:
        if p["severity"][form] == "不同意":
            sev = "安定" if sev.startswith("要注意") else "要注意（念のため）"
        lines.append(f"病状の重さ：{sev}")
    for field, numbers in (("患者の要約", (2, 3)), ("すること", (4, 5)), ("気をつけること", (6,))):
        texts = []
        for no in numbers:
            if (form, no) in p["disputes"]:
                text = p["disputes"][(form, no)][3]
            elif no in p["omit"][form]:
                continue
            else:
                text = p["items"][no]
            if form == "c" and no in SOURCE_TIMES:
                text += f"〔電子カルテ {date} {SOURCE_TIMES[no]}〕"
            texts.append(text)
        if texts:
            lines.append(f"{field}：" + "。".join(texts))
    said = p["items"][7] if 7 not in p["omit"][form] else "内容を確かめた"
    lines.append(f"受け手の要約（受け手役が書いた）：{said}")
    return lines


def make_handoffs(path: Path, patients, date) -> None:
    """申し送りの文（形を伏せ、記号だけを付ける）"""
    doc = base_doc()
    doc.add_paragraph().add_run("申し送りの文（I-PASS。形は記号で伏せてある）").bold = True
    for p in patients:
        doc.add_paragraph().add_run(p["id"]).bold = True
        for label in LABELS:
            doc.add_paragraph().add_run(f"記号 {label}").bold = True
            for line in handoff_lines(p, p["labels"][label], date):
                doc.add_paragraph(line)
    fix_doc(doc, "申し送りの文")
    doc.save(path)


def form_line(form, values):
    fields, disagree, minutes, noticed = values
    detail = "、".join(f"{name} {n}" for name, n in zip(COUNTED, fields))
    return (f"({form}) {FORM_NAMES[form]}：漏れ {sum(fields)}件（{detail}）、見立ての不同意 {disagree}件、"
            f"1件あたり {minutes:.1f}分、受け手が気づいた欠け {noticed}件")


def make_report(path: Path) -> None:
    """前回の比較報告（示すファイル。前回の成果物）"""
    summary = summarize(PREVIOUS)
    doc = base_doc()

    def line(text: str, bold: bool = False) -> None:
        doc.add_paragraph().add_run(text).bold = bold

    line("結論：(a)（AI を入れない）", bold=True)
    line("理由：(b) は漏れの合計が (a) より4件少ないが、1件あたりの時間が2.5分増えて2分を超えた。"
         "(c) は漏れの合計が (a) より2件少ないだけで、3件に届かなかった。")
    line("形ごとの結果", bold=True)
    for form in FORMS:
        line(form_line(form, summary[form]))
    line("※ 受け手の要約の欄は、漏れに数えていない（受け手役が書く欄のため）。")
    line("※ 1件あたりの時間は、書く時間と直す時間を足した分の平均。")
    line("患者別の漏れ", bold=True)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("患者", "(a)", "(b)", "(c)")):
        cell.text = text
    for p in PREVIOUS:
        counts = [sum(1 for no in p["omit"][form] if ITEM_FIELDS[no] in COUNTED) for form in FORMS]
        cells = table.add_row().cells
        for cell, text in zip(cells, [f"{p['id']}（{p['room']}号室）", *map(str, counts)]):
            cell.text = text
    fix_doc(doc, "前回の比較報告")
    doc.save(path)


if __name__ == "__main__":
    check_data(PREVIOUS, PREVIOUS_EXPECTED, PREVIOUS_CONCLUSION)
    check_data(INNER, INNER_EXPECTED, INNER_CONCLUSION)
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(here / "files"))
    parser.add_argument("--inner", default=str(here / "inner" / "inputs"))
    args = parser.parse_args()
    out, inner = Path(args.out), Path(args.inner)
    out.mkdir(parents=True, exist_ok=True)
    inner.mkdir(parents=True, exist_ok=True)
    make_report(out / "前回の比較報告.docx")
    make_judgments(out / "前回の判定シート.xlsx", PREVIOUS)
    make_timelog(out / "前回の時間の記録.xlsx", PREVIOUS)
    make_reference(out / "前回の正解の一覧.xlsx", PREVIOUS)
    make_handoffs(out / "前回の申し送り.docx", PREVIOUS, "10/13")
    make_reference(inner / "reference.xlsx", INNER)
    make_judgments(inner / "judgments.xlsx", INNER)
    make_timelog(inner / "timelog.xlsx", INNER)
    make_handoffs(inner / "handoffs.docx", INNER, "10/25")
    print(f"作った：{out}、{inner}")
