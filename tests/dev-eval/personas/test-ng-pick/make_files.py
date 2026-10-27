"""試験 NG の最初の1件選びの模擬ユーザーが示すファイル（docx と xlsx）と、内側の EVAL の入力を作る。
使い方：<仮想環境>/bin/python make_files.py [--out 出力先のフォルダ] [--inner 内側の EVAL の入力のフォルダ]
（出力先の既定は files/ と inner/inputs/）。部品の版は ../requirements.txt。
題材は、アイデア創出スキル（make-idea）の1回目の試行で、ユーザー役が最初に選んだアイデア（試験部門の Excel 結果から最初に台へ戻す NG を選ぶ）。
人の名前、項目、数はすべて架空。
"""
import argparse
import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font

FIXED = datetime.datetime(2026, 10, 19, 9, 0, 0)  # 作るたびに中身が変わらないよう、作成日時を固定する
AUTHOR = "FW 開発課"

# 先週の試験結果（示すファイル）：項目 ID、項目名、判定、備考
WEEK_RESULTS = [
    ("T-1101", "UART 連続受信", "NG", "100回目で取りこぼし"),
    ("T-1102", "UART 1バイト受信", "OK", ""),
    ("T-1103", "UART 受信の割り込み", "NG", ""),
    ("T-1105", "UART 受信のエラー検出", "NG", "フレーミングエラーが出ない"),
    ("T-1201", "電源断後の再開", "NG", ""),
    ("T-1202", "電源断の記録", "保留", "記録用の治具が届いていない"),
    ("T-1301", "LED 点滅", "OK", ""),
    ("T-1302", "LED 消灯", "NG", ""),
]
# 試験仕様書の表（示すファイル）：項目 ID、機能、前提条件。同じ機能でも書き方が違うことがある（UART 受信／UART受信処理）
WEEK_SPEC = [
    ("T-1101", "UART 受信", "常温、連続100回"),
    ("T-1102", "UART 受信", "常温"),
    ("T-1103", "UART受信処理", "常温、割り込み有効、連続10回"),
    ("T-1105", "UART受信処理", "常温"),
    ("T-1201", "電源断の復帰", "常温、電源断1回"),
    ("T-1202", "電源断の復帰", "常温、電源断1回、記録あり"),
    ("T-1301", "LED 表示", "常温"),
    ("T-1302", "LED 表示", "常温"),
]

# 内側の EVAL の入力：最も多い機能（SPI 送信）で前提条件が最も少ない T-2203 は恒温槽が要るので外し、T-2205 を選ぶのが正しい
INNER_RESULTS = [
    ("T-2201", "SPI 連続送信", "NG", ""),
    ("T-2202", "SPI 1バイト送信", "OK", ""),
    ("T-2203", "低温での SPI 送信", "NG", "波形が崩れる"),
    ("T-2205", "SPI 送信の割り込み", "NG", ""),
    ("T-2207", "SPI 送信の DMA", "NG", "最後の1バイトが欠ける"),
    ("T-2301", "電源断後の再開", "NG", ""),
    ("T-2302", "電源断の記録", "NG", ""),
    ("T-2303", "電源断の連続", "保留", "電源の切り替え器を調整中"),
    ("T-2401", "LED 点滅", "OK", ""),
    ("T-2402", "LED 消灯", "NG", ""),
    ("T-2501", "ボタン長押し", "保留", ""),
]
INNER_SPEC = [
    ("T-2201", "SPI 送信", "常温、連続1000回、DMA なし"),
    ("T-2202", "SPI 送信", "常温"),
    ("T-2203", "SPI送信処理", "恒温槽（-40℃）"),
    ("T-2205", "SPI送信処理", "常温、割り込み有効"),
    ("T-2207", "SPI 送信", "常温、DMA 有効、連続100回"),
    ("T-2301", "電源断の復帰", "常温、電源断1回"),
    ("T-2302", "電源断の復帰", "常温、電源断1回、記録あり"),
    ("T-2303", "電源断の復帰", "常温、電源断10回"),
    ("T-2401", "LED 表示", "常温"),
    ("T-2402", "LED 表示", "常温"),
    ("T-2501", "ボタン入力", "常温、長押し3秒"),
]


def base_doc() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    return doc


def fix_props(doc: Document, title: str) -> None:
    props = doc.core_properties
    props.author = AUTHOR
    props.last_modified_by = AUTHOR
    props.title = title
    props.created = FIXED
    props.modified = FIXED


def make_results(path: Path, rows) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "結果"
    ws.append(["項目ID", "項目名", "判定", "備考"])
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCD", (10, 24, 8, 30)):
        ws.column_dimensions[col].width = width
    legend = wb.create_sheet("凡例")
    legend.append(["判定", "意味"])
    legend.append(["OK", "仕様どおりに動いた"])
    legend.append(["NG", "仕様どおりに動かなかった"])
    legend.append(["保留", "試験の準備が整わず、判定していない（NG に数えない）"])
    for cell in legend[1]:
        cell.font = Font(bold=True)
    legend.column_dimensions["A"].width = 8
    legend.column_dimensions["B"].width = 50
    wb.properties.creator = "試験課"
    wb.properties.lastModifiedBy = "試験課"
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.save(path)


def make_spec(path: Path, rows) -> None:
    doc = base_doc()
    doc.add_paragraph().add_run("ファームウェア 結合試験 仕様書（抜粋）").bold = True
    doc.add_paragraph("試験項目ごとに、確かめる機能と前提条件を示す。前提条件は「、」で区切る。")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("項目ID", "機能", "前提条件")):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    fix_props(doc, "試験仕様書")
    doc.save(path)


def make_memo(path: Path) -> None:
    """先週の候補メモ（示すファイル。前回の成果物）"""
    doc = base_doc()

    def line(text: str, bold: bool = False) -> None:
        doc.add_paragraph().add_run(text).bold = bold

    line("作成：鈴木")
    line("最初に台へ戻す：T-1105", bold=True)
    line("理由：NG がいちばん多い「UART 受信」（3件）で、前提条件がいちばん少ない（1つ）")
    line("機能ごとの NG の件数", bold=True)
    line("- UART 受信：3件")
    line("- 電源断の復帰：1件")
    line("- LED 表示：1件")
    fix_props(doc, "先週の候補メモ")
    doc.save(path)


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(here / "files"))
    parser.add_argument("--inner", default=str(here / "inner" / "inputs"))
    args = parser.parse_args()
    out, inner = Path(args.out), Path(args.inner)
    out.mkdir(parents=True, exist_ok=True)
    inner.mkdir(parents=True, exist_ok=True)
    make_results(out / "先週の試験結果.xlsx", WEEK_RESULTS)
    make_spec(out / "試験仕様書.docx", WEEK_SPEC)
    make_memo(out / "先週の候補メモ.docx")
    make_results(inner / "results.xlsx", INNER_RESULTS)
    make_spec(inner / "spec.docx", INNER_SPEC)
    print(f"作った：{out}、{inner}")
